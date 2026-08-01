import io
from datetime import date

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework.views import APIView

from .models import Charge
from .permissions import IsFinanceUser


def _filtered_charges(request):
    qs = Charge.objects.select_related('type_charge').all()
    categorie = request.query_params.get('categorie')
    date_debut = request.query_params.get('date_debut')
    date_fin = request.query_params.get('date_fin')
    if categorie:
        qs = qs.filter(categorie=categorie)
    if date_debut:
        qs = qs.filter(date_charge__gte=date_debut)
    if date_fin:
        qs = qs.filter(date_charge__lte=date_fin)
    return qs.order_by('date_charge')


def _fmt_montant(montant):
    return f'{montant:,.0f}'.replace(',', ' ')


class ExportPdfView(APIView):
    permission_classes = [IsFinanceUser]

    def get(self, request):
        charges = _filtered_charges(request)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, title='Charges VisionTech SARL')
        styles = getSampleStyleSheet()
        elements = [Paragraph('Charges — VisionTech SARL', styles['Title']), Spacer(1, 0.5 * cm)]

        data = [['Date', 'Catégorie', 'Type', 'Libellé', 'Bénéficiaire', 'Montant (FCFA)']]
        total = 0
        for charge in charges:
            data.append([
                charge.date_charge.strftime('%d/%m/%Y'),
                charge.get_categorie_display(),
                charge.type_charge.nom,
                charge.libelle,
                charge.beneficiaire,
                _fmt_montant(charge.montant),
            ])
            total += charge.montant
        data.append(['', '', '', '', 'Total', _fmt_montant(total)])

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -2), 0.25, colors.grey),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('LINEABOVE', (0, -1), (-1, -1), 1, colors.black),
        ]))
        elements.append(table)
        doc.build(elements)

        buffer.seek(0)
        response = HttpResponse(buffer.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="charges_{date.today().isoformat()}.pdf"'
        return response


class ExportExcelView(APIView):
    permission_classes = [IsFinanceUser]

    def get(self, request):
        charges = _filtered_charges(request)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Charges'
        headers = ['Date', 'Catégorie', 'Type', 'Libellé', 'Bénéficiaire', 'Référence', 'Statut', 'Montant (FCFA)']
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        total = 0
        for charge in charges:
            ws.append([
                charge.date_charge.strftime('%d/%m/%Y'),
                charge.get_categorie_display(),
                charge.type_charge.nom,
                charge.libelle,
                charge.beneficiaire,
                charge.reference_piece,
                charge.get_statut_display(),
                float(charge.montant),
            ])
            total += charge.montant
        ws.append(['', '', '', '', '', '', 'Total', float(total)])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

        widths = [12, 16, 22, 30, 20, 16, 12, 16]
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="charges_{date.today().isoformat()}.xlsx"'
        return response
